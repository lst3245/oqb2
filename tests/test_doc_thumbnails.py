"""Pure-logic tests for DOC thumbnail crop helpers. No create_app(), no Word."""
import os
import tempfile
import unittest
import zipfile

from PIL import Image

from app.word_com import (
    _compute_crop_box,
    _content_bbox,
    strip_docx_header_footer_refs,
)


class ComputeCropBoxTests(unittest.TestCase):
    def test_tight_crop_adds_pad_and_clamps(self):
        box = _compute_crop_box((1000, 800), (100, 40, 400, 200), pad=24,
                                symmetric_horizontal=False)
        self.assertEqual(box, (76, 16, 424, 224))

    def test_pad_clamped_to_image_bounds(self):
        box = _compute_crop_box((200, 100), (2, 1, 198, 99), pad=24,
                                symmetric_horizontal=False)
        self.assertEqual(box, (0, 0, 200, 100))

    def test_symmetric_horizontal_uses_smaller_margin(self):
        # 80 px left white, 20 px right white → crop 20-24? pad=8 →
        # side_crop = min(80, 20) - 8 = 12 from both sides.
        box = _compute_crop_box((200, 100), (80, 10, 180, 50), pad=8,
                                symmetric_horizontal=True)
        self.assertEqual(box, (12, 2, 188, 58))


class ContentBboxTests(unittest.TestCase):
    def test_drops_isolated_page_number(self):
        im = Image.new('RGB', (400, 600), (255, 255, 255))
        for x in range(40, 360):
            for y in range(30, 120):
                im.putpixel((x, y), (10, 10, 10))
        for x in range(190, 210):
            for y in range(570, 585):
                im.putpixel((x, y), (20, 20, 20))

        bbox = _content_bbox(im, 250)
        self.assertIsNotNone(bbox)
        left, top, right, bottom = bbox
        self.assertLess(bottom, 200)
        self.assertGreater(bottom, 110)
        self.assertLessEqual(top, 30)
        self.assertGreater(right, 350)
        self.assertLess(left, 45)

    def test_keeps_lower_block_outside_footer_zone(self):
        im = Image.new('RGB', (400, 600), (255, 255, 255))
        for x in range(40, 360):
            for y in range(20, 80):
                im.putpixel((x, y), (10, 10, 10))
        for x in range(40, 360):
            for y in range(250, 320):
                im.putpixel((x, y), (10, 10, 10))

        bbox = _content_bbox(im, 250)
        self.assertIsNotNone(bbox)
        _l, top, _r, bottom = bbox
        self.assertLessEqual(top, 20)
        self.assertGreaterEqual(bottom, 320)

    def test_full_page_body_is_kept(self):
        im = Image.new('RGB', (400, 600), (255, 255, 255))
        for x in range(40, 360):
            for y in range(20, 560):
                im.putpixel((x, y), (10, 10, 10))
        bbox = _content_bbox(im, 250)
        self.assertIsNotNone(bbox)
        _l, top, _r, bottom = bbox
        self.assertLessEqual(top, 20)
        self.assertGreaterEqual(bottom, 560)


class StripHeaderFooterRefsTests(unittest.TestCase):
    def test_removes_footer_ref_keeps_page_setup_and_body(self):
        from docx import Document
        from docx.shared import Twips

        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, 'src.docx')
            dst = os.path.join(tmp, 'dst.docx')
            doc = Document()
            doc.add_paragraph('Which of the following statements are positive?')
            section = doc.sections[0]
            section.page_width = Twips(11906)
            section.page_height = Twips(16838)
            section.footer.paragraphs[0].text = '166'
            doc.save(src)

            with zipfile.ZipFile(src) as z:
                xml = z.read('word/document.xml').decode('utf-8')
            self.assertIn('footerReference', xml)
            self.assertIn('pgSz', xml)

            strip_docx_header_footer_refs(src, dst)

            with zipfile.ZipFile(dst) as z:
                out = z.read('word/document.xml').decode('utf-8')
            self.assertNotIn('footerReference', out)
            self.assertNotIn('headerReference', out)
            self.assertIn('pgSz', out)
            self.assertIn('11906', out)
            self.assertIn('Which of the following statements are positive?', out)


if __name__ == '__main__':
    unittest.main()
