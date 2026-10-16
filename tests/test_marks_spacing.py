"""Pure tests for per-subject spacing-by-marks rules (no DB, no create_app)."""
import json
import unittest

from app.marks_spacing import (
    MarksSpacingError, dumps, load_rules, normalize_rules, resolve,
)

SAMPLE_RULES = {
    'before': {'bands': [{'max': None, 'mode': 'pages', 'count': 1},
                         {'max': 7, 'mode': 'lines', 'count': 0}]},
    'after': {'bands': [{'max': 5, 'mode': 'lines', 'count': 10},
                        {'max': 2, 'mode': 'lines', 'count': 4},
                        {'max': 9, 'mode': 'pages', 'count': 1}],
              'per_mark': {'base': 1, 'lines_per_mark': 2, 'max_lines': 40}},
}


class NormalizeTests(unittest.TestCase):
    def test_sorted_with_open_band_last(self):
        rules = normalize_rules(SAMPLE_RULES)
        self.assertEqual([b['max'] for b in rules['after']['bands']], [2, 5, 9])
        self.assertEqual([b['max'] for b in rules['before']['bands']], [7, None])

    def test_json_text_round_trip(self):
        rules = normalize_rules(json.dumps(SAMPLE_RULES))
        self.assertEqual(normalize_rules(dumps(rules)), rules)

    def test_empty_is_none(self):
        self.assertIsNone(normalize_rules(None))
        self.assertIsNone(normalize_rules({'before': {'bands': []}, 'after': {}}))
        self.assertIsNone(normalize_rules({'after': {'per_mark': {'base': 0, 'lines_per_mark': 0}}}))
        self.assertIsNone(dumps(None))

    def test_rejects(self):
        bad = [
            'not json',
            [1, 2],
            {'after': {'bands': [{'max': 2, 'mode': 'columns', 'count': 1}]}},
            {'after': {'bands': [{'max': None, 'mode': 'lines'}, {'max': None, 'mode': 'lines'}]}},
            {'after': {'bands': [{'max': 3, 'count': 1}, {'max': 3, 'count': 2}]}},
            {'after': {'bands': [{'max': 3, 'mode': 'pages', 'count': 0}]}},
            {'after': {'bands': [{'max': 3, 'mode': 'pages', 'count': 11}]}},
            {'after': {'bands': [{'max': 2.5, 'count': 1}]}},
            {'after': {'per_mark': {'base': 1, 'lines_per_mark': 'x'}}},
        ]
        for raw in bad:
            with self.assertRaises(MarksSpacingError, msg=repr(raw)):
                normalize_rules(raw)

    def test_load_rules_never_raises(self):
        self.assertIsNone(load_rules('{broken'))
        self.assertIsNotNone(load_rules(json.dumps(SAMPLE_RULES)))

    def test_after_per(self):
        from app.marks_spacing import after_per
        rules = normalize_rules(SAMPLE_RULES)
        self.assertEqual(rules['after']['per'], 'part')
        self.assertNotIn('per', rules['before'])
        self.assertEqual(after_per(rules), 'part')
        self.assertEqual(after_per(None), 'part')
        q = normalize_rules({'after': {'per': 'Question', 'bands': [{'max': None, 'count': 5}]}})
        self.assertEqual(after_per(q), 'question')
        self.assertIsNone(normalize_rules({'after': {'per': 'question'}}))
        with self.assertRaises(MarksSpacingError):
            normalize_rules({'after': {'per': 'page', 'bands': [{'max': None, 'count': 1}]}})

    def test_page_alias(self):
        rules = normalize_rules({'after': {'bands': [{'max': None, 'mode': 'page', 'count': 2}]}})
        self.assertEqual(rules['after']['bands'][0], {'max': None, 'mode': 'pages', 'count': 2})


class ResolveTests(unittest.TestCase):
    def setUp(self):
        self.rules = normalize_rules(SAMPLE_RULES)

    def test_band_order_and_inclusive_bound(self):
        self.assertEqual(resolve(self.rules, 'after', 0), ('lines', 4))
        self.assertEqual(resolve(self.rules, 'after', 2), ('lines', 4))
        self.assertEqual(resolve(self.rules, 'after', 3), ('lines', 10))
        self.assertEqual(resolve(self.rules, 'after', 9), ('pages', 1))

    def test_per_mark_fallback_when_no_band(self):
        self.assertEqual(resolve(self.rules, 'after', 12), ('lines', 25))
        self.assertEqual(resolve(self.rules, 'after', 30), ('lines', 40))

    def test_open_band(self):
        self.assertEqual(resolve(self.rules, 'before', 7), ('lines', 0))
        self.assertEqual(resolve(self.rules, 'before', 8), ('pages', 1))

    def test_multi_page(self):
        rules = normalize_rules({'after': {'bands': [{'max': None, 'mode': 'pages', 'count': 3}]}})
        self.assertEqual(resolve(rules, 'after', 15), ('pages', 3))

    def test_null_marks_and_missing_rules(self):
        self.assertIsNone(resolve(self.rules, 'after', None))
        self.assertIsNone(resolve(None, 'after', 4))
        rules = normalize_rules({'after': {'bands': [{'max': 2, 'count': 3}]}})
        self.assertIsNone(resolve(rules, 'after', 5))
        self.assertIsNone(resolve(rules, 'before', 1))


class _Doc:
    """Records what add_before_spacing / add_after_spacing emit."""

    def __init__(self):
        self.ops = []

    def add_page_break(self):
        self.ops.append('page')

    def add_paragraph(self, *a, **k):
        self.ops.append('line')


class GeneratorSpacingTests(unittest.TestCase):
    def setUp(self):
        from types import SimpleNamespace
        from app import generator
        self.g = generator
        self.NS = SimpleNamespace
        manual = {'before_mode': 'lines', 'before_lines': 0, 'after_mode': 'lines', 'after_lines': 1}
        self.cfg = {
            'mc': dict(manual),
            'cq': {'before_mode': 'marks', 'before_lines': 2, 'after_mode': 'marks', 'after_lines': 3},
            'marks_rules': {'ICT': normalize_rules(SAMPLE_RULES)},
        }
        root = self.NS(id=1, part=None, parent=None, children=[], qno_end=None,
                       subject='ICT', q_type=None, marks=None)
        self.qa = self.NS(id=2, part='a', parent=root, children=[], qno_end=None,
                          subject='ICT', q_type='CQ', marks=2)
        self.qb = self.NS(id=3, part='b', parent=root, children=[], qno_end=None,
                          subject='ICT', q_type='CQ', marks=9)
        root.children = [self.qa, self.qb]
        self.root = root

    def test_mode_parse(self):
        self.assertEqual(self.g._spacing_mode('marks'), 'marks')
        self.assertEqual(self.g._spacing_mode('PAGE'), 'page')
        self.assertEqual(self.g._spacing_mode('bogus'), 'lines')
        self.assertTrue(self.g._uses_marks_spacing(self.cfg))
        self.assertFalse(self.g._uses_marks_spacing({'mc': {}, 'cq': {'after_mode': 'page'}}))

    def test_before_on_whole_total_only_at_question_start(self):
        # total 11 > 7 -> new page before the question
        s = self.g.resolve_entry_spacing(self.qa, self.cfg, starts_question=True, marks_memo={})
        self.assertEqual((s['before_mode'], s['before_pages']), ('page', 1))
        later = self.g.resolve_entry_spacing(self.qb, self.cfg, starts_question=False)
        self.assertEqual((later['before_mode'], later['before_lines']), ('lines', 2))

    def test_after_on_leaf_marks(self):
        s = self.g.resolve_entry_spacing(self.qa, self.cfg, starts_question=False)
        self.assertEqual((s['after_mode'], s['after_lines']), ('lines', 4))
        s = self.g.resolve_entry_spacing(self.qb, self.cfg, starts_question=False)
        self.assertEqual((s['after_mode'], s['after_pages']), ('page', 1))

    def test_unknown_marks_and_no_rules_fall_back(self):
        self.qa.marks = None
        self.qb.marks = None
        s = self.g.resolve_entry_spacing(self.qa, self.cfg, starts_question=True, marks_memo={})
        self.assertEqual((s['before_mode'], s['before_lines'], s['after_lines']), ('lines', 2, 3))
        cfg = dict(self.cfg, marks_rules={})
        self.qa.marks = 2
        s = self.g.resolve_entry_spacing(self.qa, cfg, starts_question=True)
        self.assertEqual((s['before_lines'], s['after_lines']), (2, 3))

    def test_stem_entry_has_no_after(self):
        s = self.g.resolve_entry_spacing(self.root, self.cfg, stem_entry=True,
                                         starts_question=True, marks_memo={})
        self.assertEqual((s['before_mode'], s['before_pages']), ('page', 1))
        self.assertEqual((s['after_mode'], s['after_lines']), ('lines', 0))

    def test_manual_modes_unchanged(self):
        mc = self.NS(id=9, part=None, parent=None, children=[], qno_end=None,
                     subject='ICT', q_type='MC', marks=1)
        self.assertIs(self.g.resolve_entry_spacing(mc, self.cfg), self.cfg['mc'])

    def test_multi_page_breaks(self):
        doc = _Doc()
        self.assertTrue(self.g.add_after_spacing(doc, {'after_mode': 'page', 'after_pages': 3}))
        self.assertEqual(doc.ops, ['page'] * 3)
        doc = _Doc()
        self.g.add_before_spacing(doc, {'before_mode': 'page', 'before_pages': 2}, True, False)
        self.assertEqual(doc.ops, ['page'])
        doc = _Doc()
        self.g.add_before_spacing(doc, {'before_mode': 'page'}, False, False)
        self.assertEqual(doc.ops, ['page'])

    def test_style_parse(self):
        self.assertEqual(self.g._spacing_style('ruled'), 'ruled')
        self.assertEqual(self.g._spacing_style(' RULED '), 'ruled')
        for raw in (None, '', 'blank', 'dotted'):
            self.assertEqual(self.g._spacing_style(raw), 'blank')

    def test_ruled_after_lines(self):
        from docx import Document
        from docx.enum.text import WD_TAB_LEADER
        ruled = {'after_mode': 'lines', 'after_lines': 3, 'after_style': 'ruled'}
        doc = Document()
        self.assertFalse(self.g.add_after_spacing(doc, ruled, allow_ruled=True))
        paras = doc.paragraphs
        self.assertEqual(len(paras), 3)
        for p in paras:
            self.assertEqual(p.text, '\t')
            stops = list(p.paragraph_format.tab_stops)
            self.assertEqual(len(stops), 1)
            self.assertEqual(stops[0].leader, WD_TAB_LEADER.LINES)
        # Not allowed (answers section / inline answers) or blank style -> empty lines.
        for spacing, allow in ((ruled, False), (dict(ruled, after_style='blank'), True)):
            doc = Document()
            self.g.add_after_spacing(doc, spacing, allow_ruled=allow)
            self.assertEqual([p.text for p in doc.paragraphs], ['', '', ''])
        # Page breaks stay blank pages.
        doc = _Doc()
        self.g.add_after_spacing(doc, {'after_mode': 'page', 'after_pages': 2,
                                       'after_style': 'ruled'}, allow_ruled=True)
        self.assertEqual(doc.ops, ['page', 'page'])

    def test_override_parse(self):
        parse = self.g.parse_marks_rules_override
        self.assertEqual(parse('', {'ICT'}), {})
        out = parse(json.dumps({'ICT': {'after': {'per_mark': {'base': 0, 'lines_per_mark': 3}}},
                                'CHEM': None, 'MATC': SAMPLE_RULES}), {'ICT', 'CHEM'})
        self.assertEqual(set(out), {'ICT', 'CHEM'})  # MATC not in the selection
        self.assertEqual(resolve(out['ICT'], 'after', 2), ('lines', 6))
        self.assertIsNone(out['CHEM'])
        with self.assertRaises(MarksSpacingError) as cm:
            parse({'ICT': {'after': {'bands': [{'max': 'x'}]}}}, {'ICT'})
        self.assertTrue(str(cm.exception).startswith('ICT: '))
        for bad in ('not json', '[1]'):
            with self.assertRaises(MarksSpacingError):
                parse(bad, {'ICT'})

    def test_after_per_question(self):
        rules = normalize_rules({'after': {'per': 'question',
                                           'bands': [{'max': 5, 'count': 6},
                                                     {'max': None, 'mode': 'pages', 'count': 2}]}})
        cfg = dict(self.cfg, marks_rules={'ICT': rules})
        # (a) is not the last part: no gap; (b) ends the 11-mark question: 2 pages.
        a = self.g.resolve_entry_spacing(self.qa, cfg, starts_question=True,
                                         ends_question=False, marks_memo={})
        self.assertEqual((a['after_mode'], a['after_lines']), ('lines', 0))
        b = self.g.resolve_entry_spacing(self.qb, cfg, starts_question=False,
                                         ends_question=True, marks_memo={})
        self.assertEqual((b['after_mode'], b['after_pages']), ('page', 2))
        # Unknown total at the end falls back to the card's lines value.
        self.qa.marks = self.qb.marks = None
        b = self.g.resolve_entry_spacing(self.qb, cfg, starts_question=False,
                                         ends_question=True, marks_memo={})
        self.assertEqual((b['after_mode'], b['after_lines']), ('lines', 3))

    def test_override_wins_and_null_means_no_rules(self):
        cfg = dict(self.cfg, marks_rules={'ICT': None})
        s = self.g.resolve_entry_spacing(self.qa, cfg, starts_question=False)
        self.assertEqual((s['after_mode'], s['after_lines']), ('lines', 3))

    def test_resolved_marks_spacing_keeps_style(self):
        cfg = dict(self.cfg, cq=dict(self.cfg['cq'], after_style='ruled'))
        s = self.g.resolve_entry_spacing(self.qa, cfg, starts_question=False)
        self.assertEqual((s['after_lines'], s['after_style']), (4, 'ruled'))


if __name__ == '__main__':
    unittest.main()
